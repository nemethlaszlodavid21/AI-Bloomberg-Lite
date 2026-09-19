import hashlib
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


# ===================================================
# EXCEL SABLON OSZLOPOK
# ===================================================

EXCEL_OSZLOPOK = [
    "Dátum",
    "Ticker",
    "Eszköz neve",
    "Típus",
    "Mennyiség",
    "Ár",
    "Deviza",
    "Jutalék",
    "Megjegyzés",
]


# ===================================================
# EXCEL SABLON GENERÁLÁS
# ===================================================

def excel_sablon_keszites():

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Tranzakciók"

    # -----------------------------------------------
    # FEJLÉC
    # -----------------------------------------------

    for col_index, oszlop in enumerate(
        EXCEL_OSZLOPOK,
        start=1
    ):

        cell = sheet.cell(
            row=1,
            column=col_index,
            value=oszlop
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # -----------------------------------------------
    # OSZLOPSZÉLESSÉGEK
    # -----------------------------------------------

    szelessegek = {
        "A": 14,
        "B": 14,
        "C": 28,
        "D": 12,
        "E": 15,
        "F": 15,
        "G": 12,
        "H": 15,
        "I": 35,
    }

    for oszlop, szelesseg in szelessegek.items():

        sheet.column_dimensions[
            oszlop
        ].width = szelesseg

    sheet.freeze_panes = "A2"

    sheet.auto_filter.ref = (
        "A1:I1000"
    )

    # -----------------------------------------------
    # BUY / SELL LEGÖRDÜLŐ
    # -----------------------------------------------

    tipus_validation = DataValidation(
        type="list",
        formula1='"BUY,SELL"',
        allow_blank=False
    )

    sheet.add_data_validation(
        tipus_validation
    )

    tipus_validation.add(
        "D2:D1000"
    )

    # -----------------------------------------------
    # DEVIZA LEGÖRDÜLŐ
    # -----------------------------------------------

    deviza_validation = DataValidation(
        type="list",
        formula1='"HUF,USD,EUR,GBP,CHF"',
        allow_blank=False
    )

    sheet.add_data_validation(
        deviza_validation
    )

    deviza_validation.add(
        "G2:G1000"
    )

    # -----------------------------------------------
    # FORMÁTUMOK
    # -----------------------------------------------

    for row in range(
        2,
        1001
    ):

        sheet[
            f"A{row}"
        ].number_format = "yyyy-mm-dd"

        sheet[
            f"E{row}"
        ].number_format = "0.000000"

        sheet[
            f"F{row}"
        ].number_format = "0.000000"

        sheet[
            f"H{row}"
        ].number_format = "0.000000"

    # -----------------------------------------------
    # ÚTMUTATÓ
    # -----------------------------------------------

    info = workbook.create_sheet(
        "Útmutató"
    )

    utmutato = [
        [
            "AI Bloomberg Lite – Portfólió import sablon"
        ],
        [
            ""
        ],
        [
            "Minden sor egy BUY vagy SELL tranzakció."
        ],
        [
            "Dátum",
            "A tranzakció dátuma."
        ],
        [
            "Ticker",
            "Az instrumentum ticker kódja, pl. AAPL."
        ],
        [
            "Eszköz neve",
            "Az instrumentum neve, pl. Apple Inc."
        ],
        [
            "Típus",
            "BUY vagy SELL."
        ],
        [
            "Mennyiség",
            "A vásárolt vagy eladott darabszám."
        ],
        [
            "Ár",
            "Egy darab vételi vagy eladási ára."
        ],
        [
            "Deviza",
            "HUF, USD, EUR, GBP vagy CHF."
        ],
        [
            "Jutalék",
            "A tranzakció díja. Ha nincs, 0."
        ],
        [
            "Megjegyzés",
            "Opcionális."
        ],
        [
            ""
        ],
        [
            "Fontos:"
        ],
        [
            "Ne módosítsd a Tranzakciók munkalap oszlopneveit."
        ],
        [
            "A feltöltés önmagában nem módosítja a portfóliót."
        ],
        [
            "Az import csak külön megerősítés után történik meg."
        ],
    ]

    for row in utmutato:
        info.append(row)

    info["A1"].font = Font(
        bold=True,
        size=14
    )

    info["A14"].font = Font(
        bold=True
    )

    info.column_dimensions[
        "A"
    ].width = 45

    info.column_dimensions[
        "B"
    ].width = 70

    # -----------------------------------------------
    # BYTES
    # -----------------------------------------------

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return output.getvalue()


# ===================================================
# IMPORT ID
# ===================================================

def _excel_import_id(
    datum,
    ticker,
    tipus,
    mennyiseg,
    ar,
    deviza,
    jutalek
):

    alap = (
        f"{datum}|"
        f"{ticker}|"
        f"{tipus}|"
        f"{mennyiseg:.10f}|"
        f"{ar:.10f}|"
        f"{deviza}|"
        f"{jutalek:.10f}"
    )

    return hashlib.sha256(
        alap.encode(
            "utf-8"
        )
    ).hexdigest()


# ===================================================
# EXCEL FELDOLGOZÁS
# ===================================================

def excel_tranzakciok_feldolgozasa(
    fajl
):

    try:

        dataframe = pd.read_excel(
            fajl,
            sheet_name="Tranzakciók"
        )

    except Exception as hiba:

        raise ValueError(
            "Az Excel-fájl nem olvasható, vagy hiányzik "
            "a 'Tranzakciók' munkalap."
        ) from hiba

    # -----------------------------------------------
    # ÜRES SOROK
    # -----------------------------------------------

    dataframe = (
        dataframe
        .dropna(
            how="all"
        )
        .copy()
    )

    if dataframe.empty:

        raise ValueError(
            "A feltöltött Excel nem tartalmaz tranzakciót."
        )

    # -----------------------------------------------
    # OSZLOPOK
    # -----------------------------------------------

    hianyzo = [
        oszlop
        for oszlop in EXCEL_OSZLOPOK
        if oszlop not in dataframe.columns
    ]

    if hianyzo:

        raise ValueError(
            "Hiányzó Excel oszlopok: "
            + ", ".join(
                hianyzo
            )
        )

    dataframe = dataframe[
        EXCEL_OSZLOPOK
    ].copy()

    # -----------------------------------------------
    # NORMALIZÁLÁS
    # -----------------------------------------------

    dataframe[
        "Dátum"
    ] = pd.to_datetime(
        dataframe[
            "Dátum"
        ],
        errors="coerce"
    )

    if dataframe[
        "Dátum"
    ].isna().any():

        raise ValueError(
            "Legalább egy tranzakció dátuma hibás."
        )

    for oszlop in [
        "Mennyiség",
        "Ár",
        "Jutalék"
    ]:

        dataframe[
            oszlop
        ] = pd.to_numeric(
            dataframe[
                oszlop
            ],
            errors="coerce"
        )

        if dataframe[
            oszlop
        ].isna().any():

            raise ValueError(
                f"A(z) '{oszlop}' oszlopban hibás szám található."
            )

    dataframe[
        "Ticker"
    ] = (
        dataframe[
            "Ticker"
        ]
        .astype(str)
        .str.strip()
        .str.upper()
    )

    dataframe[
        "Eszköz neve"
    ] = (
        dataframe[
            "Eszköz neve"
        ]
        .astype(str)
        .str.strip()
    )

    dataframe[
        "Típus"
    ] = (
        dataframe[
            "Típus"
        ]
        .astype(str)
        .str.strip()
        .str.upper()
    )

    dataframe[
        "Deviza"
    ] = (
        dataframe[
            "Deviza"
        ]
        .astype(str)
        .str.strip()
        .str.upper()
    )

    dataframe[
        "Megjegyzés"
    ] = (
        dataframe[
            "Megjegyzés"
        ]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    # -----------------------------------------------
    # VALIDÁLÁS
    # -----------------------------------------------

    if (
        ~dataframe[
            "Típus"
        ].isin(
            [
                "BUY",
                "SELL"
            ]
        )
    ).any():

        raise ValueError(
            "A Típus oszlop csak BUY vagy SELL lehet."
        )

    if (
        dataframe[
            "Mennyiség"
        ] <= 0
    ).any():

        raise ValueError(
            "A mennyiség minden tranzakciónál "
            "0-nál nagyobb kell legyen."
        )

    if (
        dataframe[
            "Ár"
        ] <= 0
    ).any():

        raise ValueError(
            "Az ár minden tranzakciónál "
            "0-nál nagyobb kell legyen."
        )

    if (
        dataframe[
            "Jutalék"
        ] < 0
    ).any():

        raise ValueError(
            "A jutalék nem lehet negatív."
        )

    if (
        dataframe[
            "Ticker"
        ] == ""
    ).any():

        raise ValueError(
            "Minden tranzakcióhoz ticker szükséges."
        )

    if (
        dataframe[
            "Eszköz neve"
        ] == ""
    ).any():

        raise ValueError(
            "Minden tranzakcióhoz eszköznév szükséges."
        )

    # -----------------------------------------------
    # CANONICAL IMPORT DATAFRAME
    # -----------------------------------------------

    eredmeny = pd.DataFrame()

    eredmeny[
        "datum"
    ] = (
        dataframe[
            "Dátum"
        ]
        .dt.strftime(
            "%Y-%m-%d"
        )
    )

    eredmeny[
        "datum_ido"
    ] = None

    eredmeny[
        "ticker"
    ] = dataframe[
        "Ticker"
    ]

    eredmeny[
        "eszkoz_nev"
    ] = dataframe[
        "Eszköz neve"
    ]

    eredmeny[
        "tipus"
    ] = dataframe[
        "Típus"
    ]

    eredmeny[
        "mennyiseg"
    ] = dataframe[
        "Mennyiség"
    ].astype(float)

    eredmeny[
        "ar"
    ] = dataframe[
        "Ár"
    ].astype(float)

    eredmeny[
        "deviza"
    ] = dataframe[
        "Deviza"
    ]

    eredmeny[
        "jutalek"
    ] = dataframe[
        "Jutalék"
    ].astype(float)

    eredmeny[
        "megjegyzes"
    ] = dataframe[
        "Megjegyzés"
    ]

    eredmeny[
        "forras"
    ] = "EXCEL"

    eredmeny[
        "import_id"
    ] = [
        _excel_import_id(
            datum=sor["datum"],
            ticker=sor["ticker"],
            tipus=sor["tipus"],
            mennyiseg=sor["mennyiseg"],
            ar=sor["ar"],
            deviza=sor["deviza"],
            jutalek=sor["jutalek"]
        )
        for _, sor
        in eredmeny.iterrows()
    ]

    return eredmeny